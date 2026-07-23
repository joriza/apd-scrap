"""
Exportador de datos a formato Excel.

Este módulo proporciona la funcionalidad principal para exportar datos de las
tablas ofertas y postulantes a archivos Excel con formateo profesional y
manejo de grandes volúmenes de datos.
"""

import os
import re
import logging
from datetime import datetime
from typing import Optional, Dict, Any, List
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

# Patrón para detectar caracteres ilegales en Excel (caracteres de control)
# Excel prohíbe caracteres de control 0x00-0x08, 0x0B-0x0C, 0x0E-0x1F excepto tab(0x09), newline(0x0A), CR(0x0D)
_ILLEGAL_CHARACTERS_RE = re.compile(
    r'[\x00-\x08\x0B-\x0C\x0E-\x1F]'
)


def _sanitize_value(value: Any) -> Any:
    """
    Sanitiza un valor para que sea válido en una celda de Excel.
    
    Elimina caracteres de control ilegales que openpyxl rechaza.
    Los valores no-string se devuelven sin modificar.
    
    Args:
        value: Valor a sanitizar
        
    Returns:
        Valor sanitizado (str limpio si era string, valor original si no)
    """
    if isinstance(value, str):
        return _ILLEGAL_CHARACTERS_RE.sub('', value)
    return value

from .formatters import ExcelFormatter


class ExportError(Exception):
    """Excepción personalizada para errores de exportación."""
    pass


class ExcelExporter:
    """
    Clase principal para exportación de datos a formato Excel.
    
    Maneja la exportación de las tablas 'ofertas' y 'postulantes' con
    paginación para grandes volúmenes, formateo profesional y validación.
    """
    
    def __init__(self, db_connection):
        """
        Inicializa el exportador de Excel.
        
        Args:
            db_connection: Conexión a la base de datos
        """
        self.db = db_connection
        self.logger = logging.getLogger(__name__)
        self.formatter = ExcelFormatter()
        
        # Configuración de paginación
        self.batch_size = 10000
        
    def export_data(self, export_type: str = "ofertas", output_path: Optional[str] = None) -> str:
        """
        Exporta datos a formato Excel.
        
        Args:
            export_type: Tipo de datos a exportar ("ofertas", "postulantes", "both")
            output_path: Ruta personalizada para el archivo de salida
            
        Returns:
            str: Ruta del archivo Excel generado
            
        Raises:
            ExportError: Si ocurre un error durante la exportación
        """
        try:
            # Validar parámetros
            self._validate_export_params(export_type, output_path)
            
            # Determinar ruta de salida
            output_path = self._get_output_path(output_path, export_type)
            
            # Validar condiciones de seguridad
            self._validate_export_safety(output_path)
            
            # Crear libro de trabajo
            wb = Workbook()
            
            # Eliminar hoja predeterminada
            default_sheet = wb.active
            wb.remove(default_sheet)
            
            # Exportar según el tipo solicitado
            if export_type in ["ofertas", "both"]:
                self._export_ofertas_sheet(wb)
                
            if export_type in ["postulantes", "both"]:
                self._export_postulantes_sheet(wb)
                
            if export_type == "both":
                self._create_summary_sheet(wb)
                
            self._create_metadata_sheet(wb, export_type)
            
            # Guardar archivo
            wb.save(output_path)
            
            self.logger.info(f"Exportación completada: {output_path}")
            return output_path
            
        except Exception as e:
            self.logger.error(f"Error durante exportación: {str(e)}")
            raise ExportError(f"Falló la exportación: {str(e)}")
    
    def _validate_export_params(self, export_type: str, output_path: Optional[str]) -> None:
        """Valida los parámetros de exportación."""
        valid_types = ["ofertas", "postulantes", "both"]
        if export_type not in valid_types:
            raise ExportError(f"Tipo de exportación inválido: {export_type}. Debe ser uno de: {valid_types}")
    
    def _get_output_path(self, output_path: Optional[str], export_type: str) -> str:
        """
        Determina la ruta de salida del archivo con sufijo de fecha/hora ISO 8601.
        
        El sufijo se aplica siempre:
        - Si NO se especifica output_path: se genera nombre base `export_{tipo}`
        - Si se especifica output_path: se inserta el sufijo antes de la extensión
        
        Formato ISO 8601 filesystem-safe: `2026-07-23T13-45-30`
        (los `:` son ilegales en Windows, se reemplazan por `-`; el `T` es el
        separador fecha/hora estándar ISO 8601).
        """
        # Timestamp ISO 8601 seguro para filesystem (sin microsegundos, sin ':')
        timestamp = datetime.now().replace(microsecond=0).isoformat().replace(":", "-")
        
        if output_path:
            # Insertar sufijo antes de la extensión, preservando el directorio
            path = Path(output_path)
            new_name = f"{path.stem}_{timestamp}{path.suffix}"
            return str(path.with_name(new_name))
        
        # Nombre autogenerado en el directorio actual
        filename = f"export_{export_type}_{timestamp}.xlsx"
        return str(Path.cwd() / filename)
    
    def _validate_export_safety(self, output_path: str) -> None:
        """Valida condiciones de seguridad para la exportación."""
        # Verificar espacio en disco
        estimated_size = self._estimate_export_size()
        disk_space = self._get_available_disk_space()
        
        if estimated_size > disk_space * 0.8:
            raise ExportError("Espacio insuficiente en disco para la exportación")
        
        # Verificar permisos de escritura
        output_dir = os.path.dirname(output_path) or os.getcwd()
        if not os.access(output_dir, os.W_OK):
            raise ExportError(f"Permisos insuficientes en el directorio: {output_dir}")
    
    def _estimate_export_size(self) -> int:
        """Estima el tamaño del archivo Excel en bytes."""
        # Estimación aproximada: 1KB por registro
        count_ofertas = self.db.get_count("ofertas")
        count_postulantes = self.db.get_count("postulantes")
        total_records = count_ofertas + count_postulantes
        return total_records * 1024  # 1KB por registro
    
    def _get_available_disk_space(self) -> int:
        """Obtiene espacio disponible en disco en bytes."""
        try:
            # Para Windows
            if os.name == 'nt':
                import ctypes
                free_bytes = ctypes.c_ulonglong(0)
                ctypes.windll.kernel32.GetDiskFreeSpaceExW(
                    ctypes.c_wchar_p(os.path.dirname(os.path.abspath(__file__))),
                    None, None, ctypes.pointer(free_bytes)
                )
                return free_bytes.value
            else:
                # Para Linux/Mac
                path = os.path.dirname(os.path.abspath(__file__))
                stat = os.statvfs(path)
                return stat.f_bavail * stat.f_frsize
        except Exception:
            # Si falla, devolver un valor grande (1GB) para permitir exportación
            return 1024 * 1024 * 1024
    
    def _export_ofertas_sheet(self, workbook: Workbook) -> None:
        """Exporta datos de la tabla ofertas a una hoja de Excel."""
        self.logger.info("Exportando datos de ofertas...")
        
        sheet = workbook.create_sheet("Ofertas")
        
        total_count = self.db.get_count("ofertas")
        if total_count == 0:
            self.logger.warning("No hay datos de ofertas para exportar")
            return
        
        current_row = 1  # Fila 1 reservada para encabezados
        headers_written = False
        exported_count = 0
        
        for offset in range(0, total_count, self.batch_size):
            batch = self.db.get_all_ofertas_paginated(offset, self.batch_size)
            if not batch:
                continue
            
            # Escribir encabezados desde el primer lote
            if not headers_written:
                headers = list(batch[0].keys())
                self._create_headers(sheet, headers)
                headers_written = True
                current_row = 2  # Los datos empiezan en la fila 2
            
            # Escribir filas de datos
            for record in batch:
                for col_idx, value in enumerate(record.values(), start=1):
                    sheet.cell(row=current_row, column=col_idx, value=_sanitize_value(value))
                current_row += 1
                exported_count += 1
            
            self.logger.info(f"Progreso ofertas: {exported_count}/{total_count}")
        
        # Formatear hoja
        self.formatter.format_headers(sheet, "Ofertas")
        self.formatter.auto_adjust_columns(sheet)
        
        # Crear tabla (current_row - 1 es la última fila escrita)
        if exported_count > 0:
            self._create_excel_table(sheet, "TableOfertas", current_row - 1)
    
    def _export_postulantes_sheet(self, workbook: Workbook) -> None:
        """Exporta datos de la tabla postulantes a una hoja de Excel."""
        self.logger.info("Exportando datos de postulantes...")
        
        sheet = workbook.create_sheet("Postulantes")
        
        total_count = self.db.get_count("postulantes")
        if total_count == 0:
            self.logger.warning("No hay datos de postulantes para exportar")
            return
        
        current_row = 1  # Fila 1 reservada para encabezados
        headers_written = False
        exported_count = 0
        
        for offset in range(0, total_count, self.batch_size):
            batch = self.db.get_all_postulantes_paginated(offset, self.batch_size)
            if not batch:
                continue
            
            # Escribir encabezados desde el primer lote
            if not headers_written:
                headers = list(batch[0].keys())
                self._create_headers(sheet, headers)
                headers_written = True
                current_row = 2  # Los datos empiezan en la fila 2
            
            # Escribir filas de datos
            for record in batch:
                for col_idx, value in enumerate(record.values(), start=1):
                    sheet.cell(row=current_row, column=col_idx, value=_sanitize_value(value))
                current_row += 1
                exported_count += 1
            
            self.logger.info(f"Progreso postulantes: {exported_count}/{total_count}")
        
        # Formatear hoja
        self.formatter.format_headers(sheet, "Postulantes")
        self.formatter.auto_adjust_columns(sheet)
        
        # Crear tabla (current_row - 1 es la última fila escrita)
        if exported_count > 0:
            self._create_excel_table(sheet, "TablePostulantes", current_row - 1)
    
    def _create_summary_sheet(self, workbook: Workbook) -> None:
        """Crea una hoja de resumen con estadísticas."""
        self.logger.info("Creando hoja de resumen...")
        
        sheet = workbook.create_sheet("Resumen")
        
        # Obtener estadísticas
        count_ofertas = self.db.get_count("ofertas")
        count_postulantes = self.db.get_count("postulantes")
        
        # Datos de resumen
        summary_data = [
            ["Estadísticas de Exportación APD-Scrap", ""],
            ["", ""],
            ["Tabla Ofertas", count_ofertas],
            ["Tabla Postulantes", count_postulantes],
            ["Total Registros", count_ofertas + count_postulantes],
            ["", ""],
            ["Fecha de Exportación", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Formato", "Excel (.xlsx)"],
        ]
        
        # Escribir datos
        for row_idx, row_data in enumerate(summary_data, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                sheet.cell(row=row_idx, column=col_idx, value=value)
        
        # Formatear
        self.formatter.format_summary_headers(sheet)
        self.formatter.auto_adjust_columns(sheet)
    
    def _create_metadata_sheet(self, workbook: Workbook, export_type: str) -> None:
        """Crea una hoja con metadatos de la exportación."""
        self.logger.info("Creando hoja de metadatos...")
        
        sheet = workbook.create_sheet("Metadatos")
        
        # Metadatos
        metadata = [
            ["Información de Exportación", ""],
            ["", ""],
            ["Tipo de Exportación", export_type],
            ["Fecha Generación", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Versión APD-Scrap", "2.0.0"],
            ["Base de Datos", "SQLite"],
            ["", ""],
            ["Columnas Tabla Ofertas", 55],
            ["Columnas Tabla Postulantes", 21],
            ["", ""],
            ["Notas", "Datos exportados desde sistema APD-Scrap. Contiene información de ofertas educativas y postulaciones."]
        ]
        
        # Escribir datos
        for row_idx, row_data in enumerate(metadata, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                sheet.cell(row=row_idx, column=col_idx, value=value)
        
        # Formatear
        self.formatter.format_metadata_headers(sheet)
        self.formatter.auto_adjust_columns(sheet)
    
    def _create_headers(self, sheet, headers: List[str]) -> None:
        """Crea los encabezados en la hoja."""
        for col_idx, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=col_idx, value=header)
    
    def _create_excel_table(self, sheet, table_name: str, total_rows: int) -> None:
        """Crea una tabla de Excel en la hoja especificada."""
        from openpyxl.utils import get_column_letter
        
        # Definir rango de la tabla
        last_col_letter = get_column_letter(sheet.max_column)
        table_range = f"A1:{last_col_letter}{total_rows}"
        
        # Crear tabla
        table = Table(displayName=table_name, ref=table_range)
        
        # Estilo de tabla
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True
        )
        
        # Agregar tabla a la hoja
        sheet.add_table(table)